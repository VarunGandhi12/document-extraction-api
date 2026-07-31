"""DOCX parser backed by python-docx.

Paragraph text stays clean; heading paragraphs (by Word style) are reported
separately in `ParsedDocument.headings` so the normalizer can split sections
without any in-band markers leaking into the output.
"""

from __future__ import annotations

import io

import docx
from docx.opc.exceptions import PackageNotFoundError

from app.core.errors import ParseError
from app.parsers.base import ParsedDocument


class DocxParser:
    content_types = frozenset(
        {"application/vnd.openxmlformats-officedocument.wordprocessingml.document"}
    )

    def parse(self, data: bytes) -> ParsedDocument:
        try:
            document = docx.Document(io.BytesIO(data))
        except PackageNotFoundError as exc:
            raise ParseError("File is not a valid DOCX package.") from exc
        except Exception as exc:  # noqa: BLE001
            raise ParseError(f"Could not read DOCX: {exc}") from exc

        lines: list[str] = []
        headings: set[str] = set()
        for para in document.paragraphs:
            text = para.text.strip()
            if not text:
                lines.append("")
                continue
            style = (para.style.name or "").lower() if para.style else ""
            if style.startswith("heading") or style == "title":
                headings.add(text)
            lines.append(text)

        return ParsedDocument(
            text="\n".join(lines).strip(),
            page_count=0,
            headings=frozenset(headings),
        )
