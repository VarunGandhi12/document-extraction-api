"""Shared fixtures.

Sample documents are generated in-process so the repo carries no binary blobs
and the fixtures always match the assertions in the tests.
"""

from __future__ import annotations

import io

import docx
import fitz  # PyMuPDF
import pytest
from fastapi.testclient import TestClient

from app.main import create_app

PDF_CT = "application/pdf"
DOCX_CT = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

SAMPLE_TEXT = (
    "Jane Doe\n"
    "jane.doe@example.com | +1 415 555 0132 | https://example.com\n"
    "EXPERIENCE\n"
    "Built and shipped services from 2021 to 2026-01-15.\n"
    "EDUCATION\n"
    "B.Sc. Computer Science, 2020\n"
)


@pytest.fixture
def sample_pdf_bytes() -> bytes:
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), SAMPLE_TEXT, fontsize=11)
    data = doc.tobytes()
    doc.close()
    return data


@pytest.fixture
def sample_docx_bytes() -> bytes:
    document = docx.Document()
    document.add_heading("Jane Doe", level=1)
    document.add_paragraph("jane.doe@example.com | +1 415 555 0132 | https://example.com")
    document.add_heading("Experience", level=2)
    document.add_paragraph("Built and shipped services from 2021 to 2026-01-15.")
    document.add_heading("Education", level=2)
    document.add_paragraph("B.Sc. Computer Science, 2020")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def client() -> TestClient:
    return TestClient(create_app())
